import docx
import json
import os

doc = docx.Document('innrly-blog-fixes-final/INNRLY_Blog_Exact_Replacements_2026-10-24.docx')

pages_data = []

p_indices = []
for i, p in enumerate(doc.paragraphs):
    if p.text.startswith('Page ') and ' of 21: ' in p.text:
        p_indices.append((i, p.text))

for page_idx, (p_start, header_text) in enumerate(p_indices):
    p_end = p_indices[page_idx + 1][0] if page_idx + 1 < len(p_indices) else len(doc.paragraphs)
    slug = header_text.split(': ')[1].strip().replace('/blog/', '')
    
    # Collect paragraph text for this page
    page_paras = [doc.paragraphs[j].text.strip() for j in range(p_start, p_end) if doc.paragraphs[j].text.strip()]
    
    # Extract meta title, description, also update from paragraphs
    page_title = None
    meta_desc = None
    also_update = []
    why = None
    
    for para in page_paras:
        if para.startswith('Page title (exact):'):
            page_title = para.replace('Page title (exact):', '').strip()
        elif para.startswith('Meta description') and '(exact):' in para:
            meta_desc = para.split('(exact):', 1)[1].strip()
        elif para.startswith('Also update:'):
            also_update.append(para.replace('Also update:', '').strip())
        elif para.startswith('Why:'):
            why = para.replace('Why:', '').strip()
            
    # Table for this page is doc.tables[page_idx]
    table = doc.tables[page_idx]
    table_rows = []
    for r in table.rows[1:]:
        cells = [c.text.strip() for c in r.cells]
        if len(cells) == 3:
            table_rows.append({
                'num': cells[0],
                'find': cells[1],
                'replace': cells[2]
            })
    
    pages_data.append({
        'page_num': page_idx + 1,
        'slug': slug,
        'header': header_text,
        'why': why,
        'page_title': page_title,
        'meta_description': meta_desc,
        'also_update': also_update,
        'paragraphs': page_paras,
        'rows': table_rows
    })

os.makedirs('backend/scripts', exist_ok=True)
with open('backend/scripts/parsed_replacements.json', 'w', encoding='utf-8') as f:
    json.dump(pages_data, f, indent=2, ensure_ascii=False)

total_rows = sum(len(p['rows']) for p in pages_data)
print(f"Successfully parsed {len(pages_data)} pages with {total_rows} replacement rows.")
